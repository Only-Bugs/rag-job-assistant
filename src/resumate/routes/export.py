"""POST /export/pdf and POST /export/docx -- application package export."""

from __future__ import annotations

import io
import re
from typing import Any

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

router = APIRouter()

_SECTION_TITLES = {
    "skills": "Skills Matrix",
    "cover":  "Cover Letter",
    "emails": "Emails",
    "ats":    "ATS Summary",
}
_SECTION_ORDER = ["skills", "cover", "emails", "ats"]


class ExportContent(BaseModel):
    content: dict[str, str]


# ---------------------------------------------------------------------------
# Markdown stripping
# ---------------------------------------------------------------------------
def _strip_markdown(text: str) -> str:
    """Remove common markdown symbols, leaving readable plain text."""
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\*{1,3}([^*\n]+)\*{1,3}", r"\1", text)
    text = re.sub(r"`([^`\n]+)`", r"\1", text)
    text = re.sub(r"^[-*_]{3,}\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# PDF export (reportlab)
# ---------------------------------------------------------------------------
def _build_pdf(content: dict[str, str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontSize=14,
        fontName="Helvetica-Bold",
        spaceAfter=6,
        spaceBefore=14,
    )
    body_style = ParagraphStyle(
        "SectionBody",
        parent=styles["Normal"],
        fontSize=10,
        fontName="Helvetica",
        leading=15,
        spaceAfter=4,
    )

    story: list[Any] = []
    for i, key in enumerate(_SECTION_ORDER):
        raw = content.get(key, "").strip()
        if not raw:
            continue
        if i > 0:
            story.append(PageBreak())
        story.append(Paragraph(_SECTION_TITLES[key], heading_style))
        story.append(Spacer(1, 4 * mm))
        for line in _strip_markdown(raw).splitlines():
            stripped = line.rstrip()
            if stripped:
                safe = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, body_style))
            else:
                story.append(Spacer(1, 3 * mm))

    doc.build(story)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# DOCX export (python-docx)
# ---------------------------------------------------------------------------
def _build_docx(content: dict[str, str]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    for para in doc.paragraphs:
        p = para._element
        p.getparent().remove(p)

    for key in _SECTION_ORDER:
        raw = content.get(key, "").strip()
        if not raw:
            continue
        doc.add_heading(_SECTION_TITLES[key], level=1)
        for line in _strip_markdown(raw).splitlines():
            stripped = line.rstrip()
            para = doc.add_paragraph(stripped if stripped else "")
            para.style = "Normal"
            for run in para.runs:
                run.font.size = Pt(11)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@router.post("/export/pdf")
async def export_pdf(body: ExportContent) -> StreamingResponse:
    pdf_bytes = _build_pdf(body.content)
    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="resumate_application.pdf"'},
    )


@router.post("/export/docx")
async def export_docx(body: ExportContent) -> StreamingResponse:
    docx_bytes = _build_docx(body.content)
    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="resumate_application.docx"'},
    )
